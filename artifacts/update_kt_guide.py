from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DOCX = Path("ansible-pg-ha/POSTGRESQL_HA_KNOWLEDGE_TRANSFER_GUIDE.docx")
OUTPUT = Path("ansible-pg-ha/POSTGRESQL_HA_KNOWLEDGE_TRANSFER_GUIDE_UPDATED.docx")


def set_paragraph_text(paragraph, text):
    """Replace visible text while preserving the paragraph and its style."""
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    paragraph.add_run(text)


def set_cell_text(cell, text):
    paragraph = cell.paragraphs[0]
    set_paragraph_text(paragraph, text)
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


def find_paragraph(document, exact_text):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == exact_text:
            return paragraph
    raise RuntimeError(f"Paragraph not found: {exact_text}")


def insert_before(target, text, style):
    text = text.replace("\n+  ", "\n  ").replace("\\\\$", "\\$")
    paragraph = target.insert_paragraph_before(text)
    paragraph.style = style
    return paragraph


def add_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


document = Document(DOCX)

# Cover and document-control metadata.
set_paragraph_text(
    document.paragraphs[2],
    "UAT and Production | Architecture, Automation, Deployment, Testing, Operations and Rubrik Data Protection",
)
set_paragraph_text(
    document.paragraphs[7],
    "This guide is the primary knowledge-transfer handout for the deployed PostgreSQL HA platform. "
    "It combines the current routing architecture, repository architecture, deployment workflow, "
    "validation suite, Rubrik data-protection boundary and everyday pg_auto_failover commands for both "
    "environments. It does not replace approved change records, Rubrik procedures or the detailed recovery runbooks.",
)
set_paragraph_text(
    document.paragraphs[20],
    "RUBRIK_WAL_CUTOVER_AND_DR_RUNBOOK.md, the Rubrik PostgreSQL data-protection technical white paper, "
    "inventories, group variables, roles and templates",
)

control = document.tables[0]
set_cell_text(control.cell(4, 1), "1.1")
set_cell_text(control.cell(5, 1), "26 August 2026")

# Platform architecture: replace the retired one-to-one router/database mapping.
set_paragraph_text(
    document.paragraphs[23],
    "Each environment uses five managed virtual machines: two PostgreSQL data candidates, one "
    "pg_auto_failover monitor and two routing nodes. Production also has an external Prometheus server. "
    "Applications connect to one database VIP. Every routing node can reach both PostgreSQL candidates, "
    "and the local primary-selector accepts only the candidate that is currently writable.",
)
set_paragraph_text(
    document.paragraphs[29],
    "Figure 1. application subnet -> VIP:5432 -> local HAProxy -> local PgBouncer:6432 -> local HAProxy "
    "primary selector:6433 -> current PostgreSQL primary:5432; standby replication; monitor control path; "
    "Prometheus exporters; Rubrik protection. [Insert approved architecture diagram here.]",
)
connection_steps = {
    31: "1.  An application opens a TLS-capable PostgreSQL connection to the environment VIP on TCP 5432.",
    32: "2.  Keepalived places the VIP on exactly one routing node using unicast VRRP; the preferred node has priority 101 and its peer has priority 100.",
    33: "3.  The VIP-facing HAProxy listener forwards the session only to PgBouncer on the same routing node at 127.0.0.1:6432.",
    34: "4.  PgBouncer provides transaction pooling and sends its server connection to the local HAProxy primary selector at 127.0.0.1:6433.",
    35: "5.  The primary selector health-checks both PostgreSQL candidates with SELECT pg_is_in_recovery(); only the node returning false remains eligible, so either router can write to whichever database node pg_auto_failover has promoted.",
}
for index, text in connection_steps.items():
    set_paragraph_text(document.paragraphs[index], text)

set_paragraph_text(
    document.paragraphs[40],
    "On primary failure, the monitor promotes the healthy standby. Both routing nodes independently detect "
    "the promoted database through postgresql_primary_selector; new or re-established pooled sessions then "
    "reach the new writable node without changing the VIP or PgBouncer database definition.",
)
set_paragraph_text(
    document.paragraphs[43],
    "Both routing nodes run PgBouncer, HAProxy and Keepalived. Each router has the same two-database candidate list; there is no permanent router-to-database pairing.",
)
set_paragraph_text(
    document.paragraphs[44],
    "Keepalived moves the VIP when the current owner or its tracked local routing chain becomes unavailable. It tracks HAProxy, PgBouncer and both loopback listeners.",
)
set_paragraph_text(
    document.paragraphs[45],
    "HAProxy can bind the VIP on either router because net.ipv4.ip_nonlocal_bind=1 is enabled. The VIP listener uses local PgBouncer, while the separate loopback selector chooses the writable database.",
)
set_paragraph_text(
    document.paragraphs[46],
    "Database failover and routing failover remain independent: pg_auto_failover changes the writable database; Keepalived changes the VIP owner. HAProxy reconnects the two layers dynamically after either event.",
)
set_paragraph_text(document.paragraphs[47], "1.8 Supporting platform and data-protection services")

support_table = document.tables[8]
set_cell_text(support_table.cell(4, 1), "Rubrik RBS and RSC; selected target provider")
set_cell_text(
    support_table.cell(4, 2),
    "Physical/base backups, SecureThrift WAL ingestion, retention and PITR evidence; activation must be proven in Rubrik",
)
set_cell_text(
    support_table.cell(5, 2),
    "Retained only as an approved rollback option; never active concurrently with Rubrik",
)

environment_table = document.tables[3]
set_cell_text(environment_table.cell(7, 1), "Rubrik target; cutover evidence required")
set_cell_text(environment_table.cell(7, 2), "Rubrik target; cutover evidence required")

# Codebase and deployment wording.
set_paragraph_text(
    document.paragraphs[62],
    "5.  Create database principals, deploy the optional monitor_ssh provider only when selected, or remove "
    "legacy monitor-WAL integration and validate the Rubrik/none provider serially.",
)
set_paragraph_text(
    document.paragraphs[70],
    "The WAL cleanup role removes legacy monitor archive entries only under guarded conditions, preserves rollback keys/data, "
    "and validates active Rubrik ownership when cutover confirmation is supplied.",
)
set_paragraph_text(
    document.paragraphs[86],
    "6.  Confirm postgresql_wal_archive_provider. UAT and Production select the Rubrik target design, but "
    "rubrik_wal_cutover_confirmed=false means a full rerun remains intentionally blocked until a successful base backup, "
    "durable Rubrik archive command and WAL recovery-point evidence exist on both data candidates.",
)

roles_table = document.tables[13]
for row in roles_table.rows[1:]:
    role = row.cells[0].text.strip()
    if role == "pgbouncer":
        set_cell_text(row.cells[2], "Transaction pooling to the local HAProxy primary selector, SCRAM userlist and TLS")
    elif role == "keepalived_haproxy":
        set_cell_text(row.cells[2], "Unicast VRRP VIP, local PgBouncer ingress and dynamic primary selection across both DB candidates")
    elif role == "wal_archive_cleanup":
        set_cell_text(row.cells[2], "Removes legacy monitor archive integration; validates Rubrik/none; retains rollback keys/data")

# Add a concise Rubrik activation subsection before testing.
testing_heading = find_paragraph(document, "4. Testing instructions")
insert_before(testing_heading, "3.11 Rubrik target architecture and activation sequence", "Heading 2")
insert_before(
    testing_heading,
    "Rubrik is the intended backup authority for both environments, but the repository does not generate the "
    "vendor archive command. Rubrik discovery requires wal_level=replica and archive_mode=on; assigning the SLA "
    "must install the Rubrik archive_command. Because archive_mode is a server-start setting, enable it with a "
    "controlled secondary-first rolling procedure before declaring the cutover complete.",
    "Normal",
)
for item in [
    "Confirm exact PostgreSQL 18, Ubuntu 24.04, RSC/CDM/RBS and pg_auto_failover support with Rubrik.",
    "Install and register RBS on both data candidates; do not use the PgBouncer VIP as a physical PostgreSQL backup host.",
    "Apply source-restricted Rubrik firewall rules, then enable archive_mode=on on the current secondary and restart only that keeper.",
    "Perform a controlled switchover, enable the former primary, assign the Rubrik SLA and require a durable non-legacy archive command on both candidates.",
    "Complete an on-demand base backup, WAL switch, Rubrik recovery-point check, controlled database switchover and isolated restore/PITR test.",
]:
    insert_before(testing_heading, item, "List Bullet")
insert_before(
    testing_heading,
    "ansible db_primary:db_standby -i \"$INVENTORY\" \\\n+  -b --become-user postgres -m shell \\\n+  -a \"psql -XAtd postgres -c \\\"SELECT name || '=' || setting FROM pg_settings WHERE name IN ('wal_level','archive_mode','archive_command','archive_library','archive_timeout') ORDER BY name;\\\"\"",
    "Code Block",
)
insert_before(
    testing_heading,
    "Current-state warning: group_vars/uat.yml and group_vars/prod.yml select rubrik but retain "
    "rubrik_wal_cutover_confirmed=false. Do not describe Rubrik WAL/PITR as operational until the backup team supplies "
    "successful backup, recovery-point and restore evidence and the reviewed variables are updated.",
    "Normal",
)

acceptance_heading = find_paragraph(document, "4.7 Acceptance criteria")
set_paragraph_text(acceptance_heading, "4.8 Acceptance criteria")
insert_before(acceptance_heading, "4.7 Rubrik and database-failover validation", "Heading 2")
insert_before(
    acceptance_heading,
    "After Rubrik activation, validate protection on the current primary and repeat the validation after a controlled "
    "pg_auto_failover switchover. The promoted node must already have archive_mode=on and the Rubrik command must "
    "remain usable. Ansible output is supporting evidence; RSC job, recovery-point and restore evidence is authoritative.",
    "Normal",
)
insert_before(
    acceptance_heading,
    "ansible db_primary:db_standby -i \"$INVENTORY\" \\\n+  -b --become-user postgres -m shell \\\n+  -a \"psql -X -d postgres -P pager=off -c 'SELECT archived_count, failed_count, last_archived_wal, last_archived_time, last_failed_wal, last_failed_time FROM pg_stat_archiver;'\"",
    "Code Block",
)
for item in [
    "Force a WAL switch on the actual primary and confirm archived_count advances without repeated failures.",
    "Confirm a new Rubrik WAL/log recovery point and a successful physical/base backup in RSC.",
    "Perform an approved controlled switchover and prove Rubrik follows the promoted node.",
    "Complete an isolated Rubrik restore or PITR test before signing off the protection design.",
]:
    insert_before(acceptance_heading, item, "List Bullet")

# Testing and diagnostic expectations for dynamic routing and Rubrik. Exact
# matching keeps these edits stable even after new paragraphs are inserted.
set_paragraph_text(
    find_paragraph(document, "Validates PgBouncer, HAProxy and a SQL query through VIP -> HAProxy -> PgBouncer -> current primary."),
    "Validates PgBouncer, the VIP-facing HAProxy listener, the loopback primary selector and a SQL query through "
    "VIP -> local HAProxy -> local PgBouncer -> local primary selector -> current PostgreSQL primary.",
)
set_paragraph_text(
    find_paragraph(document, "For Rubrik, verifies a non-legacy PostgreSQL archive owner but does not prove that a Rubrik backup job completed."),
    "For Rubrik, validates PostgreSQL and firewall prerequisites plus a non-legacy archive owner after activation; "
    "it cannot prove that an RSC backup, WAL recovery point or restore completed.",
)
set_paragraph_text(
    find_paragraph(document, "Rubrik evidence to be tested by Bayshore; Ansible health output alone is not backup-success evidence."),
    "Rubrik evidence is supplied and tested by Bayshore: successful base backup, WAL/log recovery point after promotion, "
    "retention result and isolated restore/PITR. Ansible health output alone is not backup-success evidence.",
)
set_paragraph_text(
    find_paragraph(document, "Expected: One write backend is usable; the backend paired with the standby is rejected by the primary-aware check."),
    "Expected: postgresql_write shows local_pgbouncer UP; postgresql_primary_selector shows exactly one PostgreSQL "
    "candidate UP/PROCOK and the other DOWN. Both routers should agree on the same writable database.",
)
set_paragraph_text(
    find_paragraph(document, "Expected: Rubrik design: wal_level=replica, archive_mode=on and a non-legacy Rubrik command/library. /usr/local/sbin/archive-wal must not appear."),
    "Expected after completed Rubrik activation: wal_level=replica, archive_mode=on and a non-legacy Rubrik command/library "
    "on both data candidates. /usr/local/sbin/archive-wal must not appear. Before cutover, archive_mode may remain off; "
    "that is a pending target state, not successful Rubrik WAL protection.",
)

# Add routing-selector and Rubrik commands to the glossary before the appendices.
appendix_heading = find_paragraph(document, "Appendix A. Ports, paths and service quick reference")
insert_before(appendix_heading, "5.6 Dynamic primary selection and Rubrik evidence", "Heading 2")
insert_before(appendix_heading, "[READ-ONLY] Show the database selected by both routing nodes", "Normal")
insert_before(
    appendix_heading,
    "ansible routing_nodes -i \"$INVENTORY\" -b -m shell \\\n+  -a \"echo 'show stat' | socat stdio /run/haproxy/admin.sock | awk -F, '\\\$1 == \\\"postgresql_primary_selector\\\" && \\\$2 != \\\"FRONTEND\\\" && \\\$2 != \\\"BACKEND\\\" {printf \\\"%-28s status=%-6s check=%s\\\\n\\\", \\\$2, \\\$18, \\\$37}'\"",
    "Code Block",
)
insert_before(
    appendix_heading,
    "Expected: both routers show the same current primary UP/PROCOK and the standby DOWN. Investigate immediately if "
    "both candidates appear UP, neither appears UP, or the routers disagree.",
    "Normal",
)
insert_before(appendix_heading, "[READ-ONLY] Inspect PostgreSQL archiver evidence", "Normal")
insert_before(
    appendix_heading,
    "ansible db_primary:db_standby -i \"$INVENTORY\" \\\n+  -b --become-user postgres -m shell \\\n+  -a \"psql -X -d postgres -P pager=off -c 'TABLE pg_stat_archiver;'\"",
    "Code Block",
)
insert_before(
    appendix_heading,
    "Expected after Rubrik activation: the current primary records recent successful archives and failed_count does not "
    "increase repeatedly. Confirm backup and recovery-point status separately in RSC.",
    "Normal",
)

# Port and path appendices.
ports = document.tables[19]
set_cell_text(ports.cell(4, 1), "Routing loopback only")
set_cell_text(ports.cell(4, 2), "PgBouncer transaction-pool listener")
new_port_rows = [
    ("6433/tcp", "Routing loopback only", "HAProxy primary-selector listener"),
    ("12800-12801/tcp", "Rubrik sources -> data nodes", "Rubrik Backup Service listeners"),
    ("9639/tcp", "Data nodes -> Rubrik", "SecureThrift WAL/log ingestion"),
    ("111 and 32764-32769 tcp/udp", "Data nodes -> Rubrik", "Temporary NFS transport for physical backups"),
]
for values in new_port_rows:
    row = ports.add_row()
    for cell, value in zip(row.cells, values):
        set_cell_text(cell, value)
add_repeat_table_header(ports.rows[0])

paths = document.tables[20]
for values in [
    ("/pgdata/pgroot/data/postgresql.auto.conf", "Possible durable archive_mode prerequisite; coordinate ownership with Ansible and Rubrik"),
    ("/etc/rubrik/scripts/postgres/", "Rubrik-managed PostgreSQL ingestion scripts; exact command is generated by Rubrik"),
]:
    row = paths.add_row()
    for cell, value in zip(row.cells, values):
        set_cell_text(cell, value)
add_repeat_table_header(paths.rows[0])

kt_flow = document.tables[21]
set_cell_text(
    kt_flow.cell(1, 2),
    "Explain the five-hop routing chain, DB failover, routing failover, Rubrik target state and current activation gate",
)
set_cell_text(
    kt_flow.cell(6, 2),
    "Use the command glossary, routing selector, pg_stat_archiver, logs, escalation and RSC evidence process",
)

# Update core properties and request field refresh when opened in Word.
document.core_properties.title = "PostgreSQL 18 High-Availability Platform Knowledge Transfer Guide"
document.core_properties.subject = "Current UAT and Production architecture, Ansible automation, testing and Rubrik data protection"
document.core_properties.modified = datetime.now(timezone.utc)
settings = document.settings._element
update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set(qn("w:val"), "true")

document.save(OUTPUT)
print(OUTPUT.resolve())
